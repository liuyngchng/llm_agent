#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
知识库管理器 — 对标 Go 版本 internal/kb/manager.go + extract.go
管理知识库 CRUD、文件上传处理、文本提取、向量检索。
"""
import hashlib
import logging
import os
import re
import shutil
import threading
import time
from pathlib import Path
from typing import Optional

from apps.csm.embedding_client import EmbeddingClient
from apps.csm.vdb_store import LocalVectorStore, VDB_DIR

logger = logging.getLogger(__name__)

UPLOAD_DIR = "./upload_doc"
FILE_POLL_INTERVAL = 5


class KBManager:
    """知识库管理器"""

    def __init__(self, cfg: dict, store):
        self.cfg = cfg
        self.store = store
        self._emb_client = None
        self._stores = {}  # vdb_id -> LocalVectorStore
        self._stores_lock = threading.Lock()
        self._stop_ch = threading.Event()

    def _get_emb_client(self) -> EmbeddingClient:
        if not self._emb_client:
            api = self.cfg["api"]
            self._emb_client = EmbeddingClient(
                api["embedding_api_uri"],
                api["embedding_api_key"],
                api["embedding_model_name"],
            )
        return self._emb_client

    # ============================================================
    # 知识库 CRUD
    # ============================================================

    def create_kb(self, name: str, uid: str, is_public: bool = False) -> int:
        exists = self.store.check_vdb_name_exists(name, uid)
        if exists:
            raise ValueError(f"知识库名称已存在: {name}")

        id = self.store.create_vdb(name, uid, is_public)

        # 初始化向量存储
        vs = self._get_or_create_store(id)

        # 探测 embedding 维度并初始化 collection
        dim = self._get_emb_client().dimension()
        vs.ensure_collection(dim)

        return id

    def delete_kb(self, id: int, uid: str):
        vdb_info = self.store.get_vdb_by_id(id)
        if not vdb_info or vdb_info["uid"] != uid:
            raise ValueError("无权删除该知识库")

        with self._stores_lock:
            if id in self._stores:
                self._stores[id].purge()
                self._stores[id].close()
                del self._stores[id]

        # 删除文件记录中的文件
        files = self.store.get_files_by_vdb_id(id)
        for f in files:
            try:
                os.remove(f["file_path"])
            except OSError:
                pass

        self.store.delete_vdb(id)

    def get_user_kbs(self, uid: str) -> list:
        return self.store.get_user_vdbs(uid)

    def get_public_kbs(self, uid: str) -> list:
        return self.store.get_public_vdbs(uid)

    def set_default_kb(self, id: int, uid: str):
        self.store.set_default_vdb(id, uid)

    # ============================================================
    # 文件管理
    # ============================================================

    def upload_file(self, vdb_id: int, uid: str, file_name: str, file_data: bytes) -> dict:
        # 检查知识库是否存在
        vdb_info = self.store.get_vdb_by_id(vdb_id)
        if not vdb_info or vdb_info["uid"] != uid:
            raise ValueError("知识库不存在")

        os.makedirs(UPLOAD_DIR, exist_ok=True)

        task_id = str(int(time.time() * 1000000))
        saved_name = f"{task_id}_{file_name}"
        saved_path = os.path.join(UPLOAD_DIR, saved_name)

        # 计算 MD5
        file_md5 = hashlib.md5(file_data).hexdigest()

        # 保存文件
        with open(saved_path, "wb") as f:
            f.write(file_data)

        # 检查重复
        existing = self.store.check_file_md5_exists(vdb_id, file_md5)
        if existing:
            self.delete_file(existing["id"], uid)

        # 创建数据库记录
        file_id = self.store.create_file_info(
            name=file_name,
            uid=uid,
            vdb_id=vdb_id,
            task_id=task_id,
            file_path=saved_path,
            file_md5=file_md5,
        )

        return {
            "id": file_id,
            "name": file_name,
            "uid": uid,
            "vdb_id": vdb_id,
            "task_id": task_id,
            "file_path": saved_path,
            "percent": 0,
            "process_info": "文件已上传，等待处理",
            "file_md5": file_md5,
        }

    def get_files(self, vdb_id: int) -> list:
        return self.store.get_files_by_vdb_id(vdb_id)

    def delete_file(self, file_id: int, uid: str):
        finfo = self.store.get_file_by_id(file_id)
        if not finfo or finfo["uid"] != uid:
            raise ValueError("文件不存在")

        # 从向量库删除
        abs_path = os.path.abspath(finfo["file_path"])
        self._delete_vectors_by_source(finfo["vdb_id"], abs_path)

        # 删除文件
        try:
            os.remove(finfo["file_path"])
        except OSError:
            pass

        self.store.delete_file(file_id)

    # ============================================================
    # 检索
    # ============================================================

    def search_in_kb(self, query: str, vdb_id: int, uid: str,
                     top_k: int = None, score_threshold: float = None) -> str:
        if top_k is None:
            top_k = self.cfg["kb"].get("top_k", 3)
        if score_threshold is None:
            score_threshold = self.cfg["kb"].get("score_threshold", 0.1)

        vs = self._get_or_create_store(vdb_id)

        # 计算 query 向量
        query_vec = self._get_emb_client().embed_single(query)

        # 确定检索条数
        retrieve_n = top_k
        use_rerank = self.cfg["kb"].get("rerank_enabled", False) and \
                     self.cfg["api"].get("rerank_api_uri") and \
                     self.cfg["api"].get("rerank_model_name")
        if use_rerank:
            retrieve_n = self.cfg["kb"].get("rerank_retrieve_n", 15)
            if retrieve_n <= top_k:
                retrieve_n = top_k * 3
            if retrieve_n > 50:
                retrieve_n = 50

        results = vs.search(query_vec, retrieve_n, score_threshold)

        # Rerank 重排序（简化版，暂不实现完整 rerank 客户端）
        parts = []
        for r in results[:top_k]:
            content = r["content"].replace("\n", "")
            if "......................." in content:
                continue
            parts.append(content)

        return "\n".join(parts)

    def search_all_kbs(self, query: str, uid: str,
                       top_k: int = None, score_threshold: float = None) -> str:
        if top_k is None:
            top_k = self.cfg["kb"].get("top_k", 3)
        if score_threshold is None:
            score_threshold = self.cfg["kb"].get("score_threshold", 0.1)

        kb_list = self.store.get_user_vdbs(uid)
        all_parts = []

        for kb in kb_list:
            ctx = self.search_in_kb(query, kb["id"], uid, top_k, score_threshold)
            if ctx:
                all_parts.append(f"[{kb['name']}]\n{ctx}")

        return "\n".join(all_parts)

    # ============================================================
    # 文档处理 Worker
    # ============================================================

    def start_file_worker(self):
        """启动后台文件处理线程"""
        t = threading.Thread(target=self._file_worker_loop, daemon=True, name="file_worker")
        t.start()
        logger.info("文件处理 worker 已启动")

    def stop_file_worker(self):
        self._stop_ch.set()
        logger.info("文件处理 worker 已停止")

    def _file_worker_loop(self):
        while not self._stop_ch.is_set():
            self._process_pending_files()
            self._stop_ch.wait(FILE_POLL_INTERVAL)

    def _process_pending_files(self):
        files = self.store.get_unprocessed_files()
        for f in files:
            try:
                self._process_file(f)
            except Exception as e:
                logger.error("处理文件失败: name=%s, error=%s", f["name"], e)
                self.store.update_file_progress(f["id"], 0, f"处理失败: {e}")

    def _process_file(self, finfo: dict):
        logger.info("开始处理文件: name=%s, id=%s", finfo["name"], finfo["id"])
        self.store.update_file_progress(finfo["id"], 1, "开始处理文档")

        # 提取文本
        ext = os.path.splitext(finfo["file_path"])[1].lower()
        text = self._extract_text(finfo["file_path"], ext)
        if not text.strip():
            self.store.update_file_progress(finfo["id"], 100, "文件内容为空")
            return

        # 文本切分
        chunks = self._split_text(
            text,
            self.cfg["kb"].get("chunk_size", 300),
            self.cfg["kb"].get("chunk_overlap", 80),
        )
        if not chunks:
            self.store.update_file_progress(finfo["id"], 100, "无可切分的文本内容")
            return

        logger.info("文件已切分: name=%s, chunks=%d", finfo["name"], len(chunks))
        self.store.update_file_progress(
            finfo["id"], 5, f"已切分为 {len(chunks)} 个文本块，开始向量化"
        )

        # 初始化向量存储
        vs = self._get_or_create_store(finfo["vdb_id"])
        dim = self._get_emb_client().dimension()
        vs.ensure_collection(dim)

        # 批量向量化并插入
        batch_size = 10
        total_chunks = len(chunks)
        file_name = os.path.basename(finfo["file_path"])
        abs_path = os.path.abspath(finfo["file_path"])

        for i in range(0, total_chunks, batch_size):
            end = min(i + batch_size, total_chunks)
            batch = chunks[i:end]

            # 批量 embedding
            embeddings = self._get_emb_client().embed(batch)

            # 构建记录
            records = []
            for j, chunk_text in enumerate(batch):
                records.append({
                    "id": f"{file_name}_chunk_{i + j}",
                    "vector": embeddings[j],
                    "content": chunk_text,
                    "metadata": {"source": abs_path},
                })

            # 插入向量存储
            vs.insert(records)

            # 更新进度
            percent = end / total_chunks * 100
            if percent > 99:
                percent = 99
            self.store.update_file_progress(
                finfo["id"], percent,
                f"已处理 {end}/{total_chunks} 个文本块",
            )

        self.store.update_file_progress(
            finfo["id"], 100,
            f"处理完成，共 {total_chunks} 个文本块",
        )
        logger.info("文件处理完成: name=%s", finfo["name"])

    # ============================================================
    # 内部方法
    # ============================================================

    def _get_or_create_store(self, vdb_id: int) -> LocalVectorStore:
        with self._stores_lock:
            if vdb_id in self._stores:
                return self._stores[vdb_id]

            vs = LocalVectorStore(vdb_id)
            self._stores[vdb_id] = vs
            return vs

    def _delete_vectors_by_source(self, vdb_id: int, source: str):
        with self._stores_lock:
            vs = self._stores.get(vdb_id)
            if vs:
                vs.delete_by_source(source)

    # ============================================================
    # 文本提取
    # ============================================================

    def _extract_text(self, file_path: str, ext: str) -> str:
        """根据文件后缀提取文本"""
        # 检查缓存
        txt_path = file_path + ".txt"
        if os.path.exists(txt_path):
            with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
            if text.strip():
                return text

        if ext == ".pdf":
            text = self._extract_pdf(file_path)
        elif ext == ".docx":
            text = self._extract_docx(file_path)
        elif ext in (".xlsx", ".xls"):
            text = self._extract_xlsx(file_path)
        else:
            # txt/md 直接读
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()

        # 缓存
        if text.strip():
            try:
                with open(txt_path, "w", encoding="utf-8") as f:
                    f.write(text)
            except OSError:
                pass

        return text

    def _extract_pdf(self, file_path: str) -> str:
        """从 PDF 提取文本"""
        try:
            import PyPDF2
        except ImportError:
            raise ImportError("需要安装 PyPDF2: pip install PyPDF2")

        text_parts = []
        with open(file_path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    text_parts.append(text)

        result = "\n".join(text_parts).strip()
        if not result:
            raise ValueError("PDF 文件无文本内容，可能是扫描件")
        return result

    def _extract_docx(self, file_path: str) -> str:
        """从 DOCX 提取文本"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("需要安装 python-docx: pip install python-docx")

        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        result = "\n".join(paragraphs)
        if not result:
            raise ValueError("DOCX 文件无文本内容")
        return result

    def _extract_xlsx(self, file_path: str) -> str:
        """从 XLSX 提取文本"""
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".xls":
            raise ValueError("旧版 .xls 格式不支持，请转换为 .xlsx 格式后再上传")

        try:
            import openpyxl
        except ImportError:
            raise ImportError("需要安装 openpyxl: pip install openpyxl")

        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        rows = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append("\t".join(cells))
        wb.close()

        result = "\n".join(rows)
        if not result:
            raise ValueError("XLSX 文件无内容")
        return result

    # ============================================================
    # 文本切分
    # ============================================================

    @staticmethod
    def _split_text(text: str, chunk_size: int, chunk_overlap: int) -> list[str]:
        """简单文本切分"""
        if chunk_size <= 0:
            chunk_size = 300

        paragraphs = text.split("\n")
        chunks = []
        current = ""

        for para in paragraphs:
            para = para.strip()
            if not para:
                continue

            if len(para) <= chunk_size:
                if not current:
                    current = para
                else:
                    combined = current + "\n" + para
                    if len(combined) <= chunk_size:
                        current = combined
                    else:
                        chunks.append(current)
                        current = para
            else:
                # 长段落切分
                if current:
                    chunks.append(current)
                    current = ""
                for i in range(0, len(para), chunk_size - chunk_overlap):
                    end = i + chunk_size
                    if end > len(para):
                        end = len(para)
                    chunks.append(para[i:end])
                    if end == len(para):
                        break

        if current:
            chunks.append(current)

        return chunks